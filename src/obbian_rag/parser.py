"""Admin-only bounded parsers. No URLs, uploads, macros, OCR, or external resources are executed."""

import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from .guardrails import EMAIL, INJECTION, SECRET

MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 500_000


def parse_file(path: Path) -> list[dict]:
    if path.is_symlink() or not path.is_file() or path.stat().st_size > MAX_BYTES:
        raise ValueError("Document must be a regular file no larger than 10 MiB.")
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt", ".html"}:
        text = path.read_text(encoding="utf-8-sig")
        if suffix == ".html":
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(text, "html.parser")
            for node in soup(["script", "style", "iframe", "noscript", "form"]):
                node.decompose()
            for node in soup.select('[hidden], [aria-hidden="true"]'):
                node.decompose()
            text = soup.get_text("\n", strip=True)
        sections = [{"text": text, "locator": "document"}]
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path, strict=True)
        if reader.is_encrypted or len(reader.pages) > 100:
            raise ValueError("Encrypted PDFs and PDFs over 100 pages are not supported.")
        sections = []
        for page, content in enumerate(reader.pages, 1):
            text = content.extract_text() or ""
            if not text.strip():
                raise ValueError(f"PDF page {page} has no extractable text; run reviewed OCR separately.")
            sections.append({"text": text, "locator": f"page {page}"})
    elif suffix == ".docx":
        with zipfile.ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("DOCX expanded size exceeds 50 MiB.")
            if len(archive.infolist()) > 2000:
                raise ValueError("DOCX archive contains too many entries.")
        from docx import Document

        doc = Document(path)
        sections = [
            {"text": p.text, "locator": f"paragraph {i}"}
            for i, p in enumerate(doc.paragraphs, 1)
            if p.text.strip()
        ]
        for i, table in enumerate(doc.tables, 1):
            # Repeat headers per row so isolated chunks retain the meaning of values.
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            for n, row in enumerate(rows[1:], 2):
                sections.append(
                    {
                        "text": "; ".join(f"{k}: {v}" for k, v in zip(rows[0], row)),
                        "locator": f"table {i}, row {n}",
                    }
                )
    else:
        raise ValueError("Supported formats: .md, .txt, .html, .pdf, .docx")
    if not sections or sum(len(s["text"]) for s in sections) > MAX_TEXT:
        raise ValueError("Document is empty or extracted text exceeds 500,000 characters.")
    for section in sections:
        section["text"] = section["text"].replace("\x00", "").strip()
        if (
            INJECTION.search(section["text"])
            or SECRET.search(section["text"])
            or EMAIL.search(section["text"])
        ):
            raise ValueError(
                "Document contains suspicious instructions, credentials, or personal contact data. Review before indexing."
            )
    return sections


def parse_isolated(path: Path) -> list[dict]:
    result = subprocess.run(
        [sys.executable, "-m", "obbian_rag.parser", str(path.resolve())],
        capture_output=True,
        text=True,
        timeout=20,
        check=False,
    )
    if result.returncode:
        raise ValueError("Document parsing failed: " + result.stderr[-300:])
    return json.loads(result.stdout)


def chunk_sections(sections: list[dict], max_chars: int = 1800) -> list[dict]:
    chunks = []
    for section in sections:
        blocks = re.split(r"\n\s*\n", section["text"])
        heading = section["locator"]
        for block in blocks:
            block = block.strip()
            if not block:
                continue
            if block.startswith("#") and "\n" not in block:
                heading = block.lstrip("# ").strip()
                continue
            # Oversized text split on sentence boundaries; never drop characters silently.
            pieces = re.split(r"(?<=[.!?])\s+", block) if len(block) > max_chars else [block]
            pending = ""
            for piece in pieces:
                if len(piece) > max_chars:
                    raise ValueError("A passage exceeds the chunk limit; add reviewed paragraph boundaries.")
                if pending and len(pending) + len(piece) + 1 > max_chars:
                    chunks.append({"text": pending, "locator": f"{section['locator']} / {heading}"})
                    pending = ""
                pending = (pending + " " + piece).strip()
            if pending:
                chunks.append({"text": pending, "locator": f"{section['locator']} / {heading}"})
    return chunks


if __name__ == "__main__":
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
        if sys.platform.startswith("linux"):
            resource.setrlimit(resource.RLIMIT_AS, (768 * 1024**2, 768 * 1024**2))
        print(json.dumps(parse_file(Path(sys.argv[1]))))
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
